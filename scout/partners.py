"""The client's existing creator partnerships: marking, hold-out recall, optional seeds.

The list is private client data, read from ``SCOUT_PARTNERS_FILE`` or
``private/prenew-example-collaborations.xlsx`` (gitignored). It only lists
creator names, so matching compares normalised names against handles and
display names on every platform.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .config import ROOT
from .store import Store

DEFAULT_PATH = ROOT / "private" / "prenew-example-collaborations.xlsx"


def norm(text: str | None) -> str:
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


@dataclass
class Partner:
    key: str
    market: str
    country: str
    names: set[str] = field(default_factory=set)
    platforms: set[str] = field(default_factory=set)
    niche: str = ""
    yt_subscribers: float | None = None
    tiktok_followers: float | None = None

    def matches(self, *texts: str | None) -> bool:
        for t in texts:
            n = norm(t)
            if not n:
                continue
            for name in self.names:
                if n == name:
                    return True
                # Containment (e.g. "somegamer" in "somegamertv") only when the
                # shorter name is distinctive and most of the longer one.
                short, long_ = sorted((n, name), key=len)
                if len(short) >= 6 and short in long_ and len(short) / len(long_) >= 0.6:
                    return True
        return False


def partners_path() -> Path:
    return Path(os.environ.get("SCOUT_PARTNERS_FILE", DEFAULT_PATH))


def load_partners(path: Path | None = None) -> list[Partner]:
    path = path or partners_path()
    if not path.is_file():
        return []
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows)]

    def col(row, name):
        for i, h in enumerate(header):
            if h.startswith(name):
                return row[i]
        return None

    merged: dict[tuple[str, str], Partner] = {}
    for row in rows:
        key = col(row, "creator key")
        if not key:
            continue
        market = str(col(row, "market") or "").upper()
        p = merged.setdefault((norm(key), market), Partner(key=str(key), market=market,
                                                           country=str(col(row, "country") or "")))
        channel = str(col(row, "creator / channel") or "")
        p.names |= {n for n in (norm(key), norm(channel), norm(re.sub(r"\(.*?\)", "", channel)))
                    if n}
        for part in re.findall(r"\(([^)]+)\)", channel):
            p.names.add(norm(part))
        platform = str(col(row, "platform") or "")
        p.platforms |= {x.strip() for x in re.split(r"[+/]", platform) if x.strip()}
        p.niche = p.niche or str(col(row, "niche") or "")
        p.yt_subscribers = p.yt_subscribers or col(row, "yt subscribers")
        p.tiktok_followers = p.tiktok_followers or col(row, "tiktok followers")
    return list(merged.values())


class PartnerIndex:
    def __init__(self, partners: list[Partner]):
        self.partners = partners

    def match(self, creator: dict, market: str | None = None) -> Partner | None:
        for p in self.partners:
            if market and p.market != market:
                continue
            if p.matches(creator.get("handle"), creator.get("nickname")):
                return p
        return None

    def for_market(self, market: str) -> list[Partner]:
        return [p for p in self.partners if p.market == market]


def recall_report(store: Store, market: str, run_id: int | None = None,
                  partners: list[Partner] | None = None) -> list[dict]:
    """For each partner in ``market``: found or not, and the stage that lost it."""
    partners = partners if partners is not None else load_partners()
    creators = [dict(r) for r in store.conn.execute(
        "SELECT platform, uid, handle, nickname, followers, region FROM creators").fetchall()]
    out = []
    for p in [x for x in partners if x.market == market]:
        matches = [c for c in creators if p.matches(c["handle"], c["nickname"])]
        entry = {"partner": p.key, "platforms": sorted(p.platforms), "niche": p.niche,
                 "found": False, "stage": "never reached", "accounts": []}
        best_rank = 99
        for c in matches:
            s = store.get_screening(market, c["platform"], c["uid"]) or {}
            pj = store.get_prejudgment(market, c["platform"], c["uid"]) or {}
            d = store.get_decision(market, c["platform"], c["uid"]) or {}
            status = s.get("status")
            stage, rank = _stage(status, s.get("reason"), pj, d)
            entry["accounts"].append({
                "platform": c["platform"], "handle": c["handle"], "nickname": c["nickname"],
                "followers": c["followers"], "region": c["region"], "status": status,
                "reason": s.get("reason"), "prejudge": pj.get("verdict"),
                "decision_reason": (d.get("data") or {}).get("reasons"),
                "found_via": [dict(e) for e in store.edges_to(c["platform"], c["uid"])][:3],
                "stage": stage})
            if rank < best_rank:
                best_rank, entry["stage"] = rank, stage
                # Finding a partner that only misses the size band still counts:
                # discovery worked, the band is a campaign choice.
                entry["found"] = status in ("accepted", "maybe") or \
                    (status == "filtered" and (s.get("reason") or "").startswith(("above band",
                                                                                 "below band")))
        out.append(entry)
    return out


def _stage(status: str | None, reason: str | None, pj: dict, d: dict) -> tuple[str, int]:
    if status == "accepted":
        return "found and accepted", 0
    if status == "maybe":
        return "found, judged maybe", 1
    if status == "needs_judgment":
        return "enriched, not judged yet", 2
    if status == "rejected":
        return f"rejected by the judge: {((d.get('data') or {}).get('reasons') or '')[:120]}", 3
    if status == "pending":
        return "seen, waiting for enrichment (budget or priority)", 4
    if status == "filtered" and (reason or "").startswith(("above band", "below band")):
        return f"found, outside the size band: {reason}", 2
    if status == "filtered" and (reason or "").startswith("pre-judge"):
        return f"skipped by the pre-judge: {reason}", 5
    if status == "filtered":
        return f"filtered: {reason}", 6
    if status == "other_market":
        return f"placed in another market: {reason}", 7
    if status:
        return f"{status}: {reason or ''}", 8
    return "seen in another market's data only", 9
