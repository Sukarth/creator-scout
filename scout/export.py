"""XLSX and CSV export of a run's results."""

from __future__ import annotations

import csv
import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import metrics as metrics_mod
from .store import Store

PLATFORM = "tiktok"
OUTREACH_STATUSES = ["to contact", "contacted", "replied", "deal", "declined"]

SHORTLIST_COLUMNS = [
    "status", "platform", "handle", "profile_url", "nickname", "market", "market_confidence",
    "market_evidence", "followers", "median_views", "er_views", "posts_per_week",
    "last_post_date", "decision", "fit_score", "gaming_pc_relevance", "niche_tags",
    "content_styles", "trust_content_score", "reasons", "evidence_quote", "sponsors_mentioned",
    "competitor_conflict", "emails", "other_links", "bio", "bio_link", "content_language",
    "suggested_deal", "price_low_eur", "price_high_eur", "pitch_language", "pitch_subject",
    "pitch_body", "dm_text", "found_via", "prejudge", "first_seen",
]


def profile_url(handle: str | None) -> str:
    return f"https://www.tiktok.com/@{handle}" if handle else ""


def _date(ts: float | int | None) -> str:
    return dt.datetime.fromtimestamp(ts, dt.timezone.utc).strftime("%Y-%m-%d") if ts else ""


def found_via(store: Store, uid: str, limit: int = 4) -> str:
    parts: list[str] = []
    for e in store.edges_to(PLATFORM, uid):
        label = {"hashtag": f"#{e['via']}", "keyword": f"search '{e['via']}'",
                 "following": f"followed by @{e['via']}",
                 "retailer_search": f"user search '{e['via']}'"}.get(e["kind"], f"{e['kind']}:{e['via']}")
        if label not in parts:
            parts.append(label)
    more = f" (+{len(parts) - limit} more)" if len(parts) > limit else ""
    return "; ".join(parts[:limit]) + more


def creator_row(store: Store, market: str, screening: dict) -> dict:
    uid = screening["uid"]
    c = store.get_creator(PLATFORM, uid) or {}
    m = store.get_metrics(PLATFORM, uid) or {}
    d = _decision(store, market, uid)
    pitch = store.get_pitch(market, PLATFORM, uid) or {}
    pj = store.get_prejudgment(market, PLATFORM, uid) or {}
    price = metrics_mod.price_estimate(m.get("median_views"))
    confidence = {"sure": "high", "unsure": "medium"}.get(screening.get("bucket"), "low")
    competitors = m.get("competitors") or []
    return {
        "status": "",
        "platform": PLATFORM,
        "handle": c.get("handle"),
        "profile_url": profile_url(c.get("handle")),
        "nickname": c.get("nickname"),
        "market": d.get("market_resolution") or (market if screening.get("bucket") == "sure" else
                                                  f"{market}?"),
        "market_confidence": confidence,
        "market_evidence": "; ".join(screening.get("market_evidence") or []),
        "followers": c.get("followers"),
        "median_views": m.get("median_views"),
        "er_views": m.get("er_views"),
        "posts_per_week": m.get("posts_per_week"),
        "last_post_date": _date(m.get("last_post_at")),
        "decision": d.get("decision") or ("not judged" if screening["status"] == "needs_judgment"
                                          else screening["status"]),
        "fit_score": d.get("fit_score"),
        "gaming_pc_relevance": d.get("gaming_pc_relevance"),
        "niche_tags": ", ".join(d.get("niche_tags") or []),
        "content_styles": ", ".join(d.get("content_styles") or []),
        "trust_content_score": d.get("trust_content_score"),
        "reasons": d.get("reasons"),
        "evidence_quote": d.get("evidence_quote"),
        "sponsors_mentioned": ", ".join(sorted(set((m.get("sponsors") or [])
                                                   + (d.get("sponsors_mentioned") or [])))),
        "competitor_conflict": ", ".join(competitors) if competitors else
        ("yes" if d.get("competitor_conflict") else "no"),
        "emails": ", ".join(c.get("emails") or []) or "no contact found",
        "other_links": ", ".join(f"{k}: {v}" for k, v in (c.get("links") or {}).items()),
        "bio": c.get("bio"),
        "bio_link": c.get("bio_link"),
        "content_language": d.get("content_language"),
        "suggested_deal": metrics_mod.suggest_deal(c.get("followers"), m.get("old_hardware")),
        "price_low_eur": price[0] if price else None,
        "price_high_eur": price[1] if price else None,
        "pitch_language": pitch.get("language"),
        "pitch_subject": pitch.get("subject"),
        "pitch_body": pitch.get("body"),
        "dm_text": pitch.get("dm"),
        "found_via": found_via(store, uid),
        "prejudge": f"{pj['verdict']}: {pj.get('reason') or ''}".strip(": ") if pj else "",
        "first_seen": _date(c.get("first_seen_at")),
    }


def _decision(store: Store, market: str, uid: str) -> dict:
    row = store.conn.execute(
        "SELECT decision, fit_score, data FROM decisions WHERE market = ? AND platform = ? AND uid = ?",
        (market, PLATFORM, uid),
    ).fetchone()
    if row is None:
        return {}
    data = json.loads(row["data"] or "{}")
    data.setdefault("decision", row["decision"])
    data.setdefault("fit_score", row["fit_score"])
    return data


def build_sheets(store: Store, run_id: int) -> dict[str, list[dict]]:
    run = store.get_run(run_id)
    if run is None:
        raise ValueError(f"run {run_id} not found")
    market = run["market"]
    scr = store.screenings(market, run_id=run_id)

    def rows(statuses: tuple[str, ...]) -> list[dict]:
        out = [creator_row(store, market, s) for s in scr if s["status"] in statuses]
        out.sort(key=lambda r: (-(r["fit_score"] or 0), -(r["followers"] or 0)))
        return out

    if (run["params"] or {}).get("judge", "none") == "none":
        shortlist, maybe = rows(("accepted", "needs_judgment")), rows(("maybe",))
    else:
        # Judged runs: only accepted creators are shortlisted; enriched but not yet
        # judged creators (e.g. when the budget ran out) go with the maybes.
        shortlist, maybe = rows(("accepted",)), rows(("maybe", "needs_judgment"))

    other = []
    for s in scr:
        if s["status"] != "other_market":
            continue
        c = store.get_creator(PLATFORM, s["uid"]) or {}
        other.append({"handle": c.get("handle"), "profile_url": profile_url(c.get("handle")),
                      "region": c.get("region") or "unknown", "region_source": c.get("region_source"),
                      "language": c.get("language"), "followers": c.get("followers"),
                      "bio": c.get("bio"), "found_via": found_via(store, s["uid"])})
    other.sort(key=lambda r: (r["region"], -(r["followers"] or 0)))

    everyone = []
    for s in scr:
        c = store.get_creator(PLATFORM, s["uid"]) or {}
        everyone.append({"handle": c.get("handle"), "status": s["status"], "reason": s["reason"],
                         "bucket": s["bucket"], "region": c.get("region"),
                         "followers": c.get("followers"),
                         "market_evidence": "; ".join(s.get("market_evidence") or []),
                         "found_via": found_via(store, s["uid"])})
    everyone.sort(key=lambda r: (r["status"], r["handle"] or ""))

    seeds = [{"seed": f"@{s['handle']}", "kind": s["kind"], "label": s["label"],
              "pages_fetched": s["pages_fetched"], "accounts_seen": s["accounts_seen"],
              "new_in_market": s["new_in_market"], "credits_spent": s["credits_spent"],
              "total_following": s["total_following"], "exhausted": bool(s["exhausted"]),
              "exhausted_reason": s["exhausted_reason"]}
             for s in store.seeds(market, active_only=False)]
    seeds += source_summary(store, run_id, market)

    funnel = run["funnel"] or {}
    log = [{"field": k, "value": v} for k, v in [
        ("run_id", run_id), ("market", market), ("brief", run["brief"]),
        ("params", json.dumps(run["params"])), ("status", run["status"]),
        ("budget", run["budget"]), ("credits_used", run["credits_used"]),
        ("api_calls", run["api_calls"]), ("cache_hits", run["cache_hits"]),
        ("llm_calls", run.get("llm_calls")), ("llm_tokens", run.get("llm_tokens")),
        ("accepted_by_first_source", json.dumps(accepted_by_first_source(store, run_id, market),
                                                ensure_ascii=False)),
        ("started", _iso(run["started_at"])), ("finished", _iso(run["finished_at"])),
        ("duration_s", round((run["finished_at"] or run["started_at"]) - run["started_at"], 1)),
    ]] + [{"field": f"funnel.{k}", "value": v} for k, v in funnel.items()]

    return {"Shortlist": shortlist, "Maybe": maybe, "Other markets pool": other,
            "Seeds and sources": seeds, "All screened": everyone, "Run log": log}


def source_summary(store: Store, run_id: int, market: str) -> list[dict]:
    """Per harvest source and snowball seed: accounts seen, in market, enriched, accepted."""
    rows = store.conn.execute(
        "SELECT e.kind, e.via, COUNT(DISTINCT e.to_uid) AS accounts,"
        " COUNT(DISTINCT CASE WHEN s.bucket IN ('sure', 'unsure') THEN e.to_uid END) AS in_market,"
        " COUNT(DISTINCT CASE WHEN s.status IN ('needs_judgment', 'accepted', 'maybe', 'rejected')"
        "   THEN e.to_uid END) AS enriched,"
        " COUNT(DISTINCT CASE WHEN s.status = 'accepted' THEN e.to_uid END) AS accepted"
        " FROM edges e LEFT JOIN screenings s ON s.uid = e.to_uid AND s.platform = e.platform"
        " AND s.market = ? WHERE e.run_id = ? AND e.kind IN ('hashtag', 'keyword', 'following')"
        " GROUP BY e.kind, e.via ORDER BY accepted DESC, in_market DESC",
        (market, run_id),
    ).fetchall()
    return [{"seed": f"{r['kind']}: {r['via']}", "kind": r["kind"], "label": "",
             "accounts_seen": r["accounts"], "new_in_market": r["in_market"],
             "enriched": r["enriched"], "accepted": r["accepted"]} for r in rows]


def accepted_by_first_source(store: Store, run_id: int, market: str) -> dict[str, int]:
    """Attribute each accepted creator to the source type that found it first."""
    out: dict[str, int] = {}
    for s in store.screenings(market, "accepted", run_id=run_id):
        edges = [e for e in store.edges_to(PLATFORM, s["uid"]) if e["run_id"] == run_id]
        if not edges:
            continue
        first = edges[0]
        label = {"hashtag": f"#{first['via']}", "keyword": f"search '{first['via']}'",
                 "following": "snowball"}.get(first["kind"], first["kind"])
        out[label] = out.get(label, 0) + 1
    return dict(sorted(out.items(), key=lambda kv: -kv[1]))


def _iso(ts: float | None) -> str:
    return dt.datetime.fromtimestamp(ts).isoformat(timespec="seconds") if ts else ""


def write_xlsx(sheets: dict[str, list[dict]], path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F2937")
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        columns = SHORTLIST_COLUMNS if name in ("Shortlist", "Maybe") else _columns(rows)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for r in rows:
            ws.append([_cell(r.get(col)) for col in columns])
        ws.freeze_panes = "B2" if name in ("Shortlist", "Maybe") else "A2"
        for i, col in enumerate(columns, start=1):
            width = min(max([len(str(col))] + [len(str(r.get(col) or ""))
                                                for r in rows[:200]]) + 2, 60)
            ws.column_dimensions[get_column_letter(i)].width = width
        if name in ("Shortlist", "Maybe") and rows:
            dv = DataValidation(type="list", formula1='"' + ",".join(OUTREACH_STATUSES) + '"',
                                allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"A2:A{len(rows) + 1}")
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=False)
            ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_csvs(sheets: dict[str, list[dict]], directory: Path) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, rows in sheets.items():
        columns = SHORTLIST_COLUMNS if name in ("Shortlist", "Maybe") else _columns(rows)
        path = directory / (name.lower().replace(" ", "_") + ".csv")
        with path.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            for r in rows:
                writer.writerow({k: _cell(r.get(k)) for k in columns})
        paths.append(path)
    return paths


def _columns(rows: list[dict]) -> list[str]:
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols or ["(empty)"]


def _cell(value):
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, bool):
        return "yes" if value else "no"
    return value


def export_run(store: Store, run_id: int, out_dir: Path, formats: tuple[str, ...] = ("xlsx", "csv")) -> list[Path]:
    run = store.get_run(run_id)
    sheets = build_sheets(store, run_id)
    stamp = dt.datetime.fromtimestamp(run["started_at"]).strftime("%Y%m%d")
    base = f"{run['market'].lower()}_{stamp}_run{run_id}"
    paths: list[Path] = []
    if "xlsx" in formats:
        paths.append(write_xlsx(sheets, out_dir / f"{base}.xlsx"))
    if "csv" in formats:
        paths += write_csvs(sheets, out_dir / base)
    return paths
